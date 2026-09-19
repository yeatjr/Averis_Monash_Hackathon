"""Multi-format document extraction and parsing module for SDOC hackathon.

Supports:
- Text files (.txt)
- PDF documents (.pdf, including scan/corruption detection)
- Word documents (.docx)
- Excel spreadsheets (.xlsx)

Handles synonym normalization for the 7 comparison fields and detects
reliability issues (unreadable files, wrong document types, missing values).
"""
import os
import re
from typing import Dict, Any, Optional, Tuple
import pypdf
import docx
import openpyxl

BLANK_TOKENS = {"???", "_______", "tba", "tbc", "n/a", "____mt", "", "none"}

WRONG_DOC_MARKERS = [
    ("commercial_invoice", [r"commercial invoice", r"\bthis is a commercial invoice\b"]),
    ("packing_list", [r"packing list", r"\bpacking list only\b"]),
    ("coo", [r"certificate of origin", r"\bcertificate of origin\b"])
]


def identify_field(label: str) -> Optional[str]:
    """Map any label synonym or bilingual header to canonical field name."""
    lbl = label.lower().strip()
    # Check notify before consignee because of "Notify Party/Intermediate Consignee"
    if "notify" in lbl or "通知人" in lbl:
        return "notify_party"
    if "consignee" in lbl or "to the order of" in lbl or "收货人" in lbl:
        return "consignee"
    if "shipper" in lbl or "发货人" in lbl:
        return "shipper"
    if "loading" in lbl or "load port" in lbl or "pol" in lbl or "装货港" in lbl:
        return "port_of_loading"
    if "discharge" in lbl or "pod" in lbl or "卸货港" in lbl:
        return "port_of_discharge"
    if "container" in lbl or "箱数" in lbl:
        return "container_count"
    if "gross" in lbl or "毛重" in lbl or "weight" in lbl:
        return "gross_weight_kg"
    return None


def is_blank(val: Any) -> bool:
    """Check if value represents a missing/unspecified field token."""
    if val is None:
        return True
    s = str(val).strip().lower()
    if s in BLANK_TOKENS:
        return True
    if "???" in s or "___" in s:
        return True
    if re.match(r"^(tba|tbc|n/a)\b", s):
        return True
    return False


def clean_party(val: str) -> str:
    """Normalize party name by removing address lines and extra whitespace."""
    if not val:
        return ""
    # In xlsx, party name is separated by ' | '
    if " | " in val:
        val = val.split(" | ")[0]
    # In docx/pdf/txt, multi-line address starts on line 2
    lines = [ln.strip() for ln in str(val).splitlines() if ln.strip()]
    name = lines[0] if lines else str(val).strip()
    return name.strip()


def clean_port(val: str) -> str:
    """Normalize port name by stripping trailing port codes like '(SGSIN)'."""
    if not val:
        return ""
    s = str(val).strip()
    # Strip (SGSIN) or similar port codes
    s = re.sub(r"\s*\([A-Z0-9/]+\)\s*$", "", s).strip()
    return s


def clean_container_count(val: Any) -> Optional[int]:
    """Parse container count integer (e.g. '6 x 40\\'HC' -> 6)."""
    if is_blank(val):
        return None
    m = re.search(r"(\d+)", str(val))
    return int(m.group(1)) if m else None


def clean_weight(val: Any) -> Optional[int]:
    """Parse gross weight kg integer (e.g. '131,058 KG' -> 131058)."""
    if is_blank(val):
        return None
    digits = re.sub(r"[^\d]", "", str(val))
    return int(digits) if digits else None


def check_wrong_doc_type(text: str) -> Optional[str]:
    """Return reason string if document text belongs to a wrong doc type."""
    t = text.lower()
    for doc_name, markers in WRONG_DOC_MARKERS:
        for m in markers:
            if re.search(m, t):
                return doc_name
    return None


def extract_from_txt(path: str) -> Tuple[Dict[str, Any], str, Optional[str]]:
    """Extract fields and check for wrong doc type from a plain text file."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    wrong_doc = check_wrong_doc_type(content)
    if wrong_doc:
        return {}, content, wrong_doc

    fields = {}
    for line in content.splitlines():
        if ":" in line:
            lbl, val = line.split(":", 1)
            fld = identify_field(lbl)
            if fld and fld not in fields:
                fields[fld] = val.strip()

    return fields, content, None


def extract_from_docx(path: str) -> Tuple[Dict[str, Any], str, Optional[str]]:
    """Extract fields and check for wrong doc type from a Word document."""
    try:
        doc = docx.Document(path)
    except Exception:
        return {}, "", "unreadable"

    all_text = []
    fields = {}
    for p in doc.paragraphs:
        if p.text.strip():
            all_text.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 2:
                lbl = row.cells[0].text.strip()
                val = row.cells[1].text.strip()
                all_text.append(f"{lbl}: {val}")
                fld = identify_field(lbl)
                if fld and fld not in fields:
                    fields[fld] = val

    content = "\n".join(all_text)
    wrong_doc = check_wrong_doc_type(content)
    if wrong_doc:
        return {}, content, wrong_doc

    return fields, content, None


def extract_from_xlsx(path: str) -> Tuple[Dict[str, Any], str, Optional[str]]:
    """Extract fields and check for wrong doc type from an Excel spreadsheet."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
    except Exception:
        return {}, "", "unreadable"

    all_text = []
    fields = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] is not None:
            lbl = str(row[0]).strip()
            val = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            all_text.append(f"{lbl}: {val}")
            fld = identify_field(lbl)
            if fld and fld not in fields:
                fields[fld] = val

    content = "\n".join(all_text)
    wrong_doc = check_wrong_doc_type(content)
    if wrong_doc:
        return {}, content, wrong_doc

    return fields, content, None


def extract_from_pdf(path: str) -> Tuple[Dict[str, Any], str, Optional[str]]:
    """Extract fields from PDF, handling unreadable / scanned cases."""
    if os.path.getsize(path) == 0:
        return {}, "", "unreadable"

    try:
        reader = pypdf.PdfReader(path)
        pages_text = [p.extract_text() or "" for p in reader.pages]
        text = "\n".join(pages_text).strip()
    except Exception:
        return {}, "", "unreadable"

    # Scanned image-only PDF detection
    if len(text) < 30 or "SCANNED COPY - NO OCR TEXT LAYER" in text:
        return {}, text, "unreadable"

    wrong_doc = check_wrong_doc_type(text)
    if wrong_doc:
        return {}, text, wrong_doc

    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    fields = {}
    i = 0
    while i < len(lines):
        line = lines[i]
        if ":" in line:
            lbl, val = line.split(":", 1)
            fld = identify_field(lbl)
            if fld in ("container_count", "gross_weight_kg"):
                fields[fld] = val.strip()
                i += 1
                continue

        fld = identify_field(line)
        if fld:
            if fld in ("container_count", "gross_weight_kg") and ":" in line:
                fields[fld] = line.split(":", 1)[1].strip()
            elif i + 1 < len(lines) and fld not in fields:
                fields[fld] = lines[i + 1].strip()
        i += 1

    return fields, text, None


def extract_document(path: str) -> Dict[str, Any]:
    """Unified entry point to extract document content and inspect health."""
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return {"status": "UNREADABLE", "reason": "unreadable", "fields": {}, "text": ""}

    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        fields, text, failure = extract_from_txt(path)
    elif ext == ".pdf":
        fields, text, failure = extract_from_pdf(path)
    elif ext == ".docx":
        fields, text, failure = extract_from_docx(path)
    elif ext == ".xlsx":
        fields, text, failure = extract_from_xlsx(path)
    else:
        return {"status": "UNREADABLE", "reason": "unreadable", "fields": {}, "text": ""}

    if failure == "unreadable":
        return {"status": "UNREADABLE", "reason": "unreadable", "fields": {}, "text": text}
    if failure:
        return {"status": "WRONG_DOC_TYPE", "reason": "wrong_doc_type", "fields": {}, "text": text}

    return {"status": "OK", "reason": None, "fields": fields, "text": text}
